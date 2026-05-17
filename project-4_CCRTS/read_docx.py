from docx import Document
paths = [
r"Phase 1 - Customer Complaint & Resolution Tracking System.docx",
r"Instructions.docx"
]
for p in paths:
    print('\n---FILE: {}'.format(p))
    try:
        doc = Document(p)
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
        for table in doc.tables:
            for row in table.rows:
                texts = [cell.text.strip() for cell in row.cells]
                print(' | '.join(texts))
    except Exception as e:
        print('ERROR reading {}: {}'.format(p, e))
